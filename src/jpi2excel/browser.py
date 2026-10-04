# Copyright © 2025-2026 by Alan M. Marcum, Nescorna Professional.
# All rights reserved.
"""Revision: 1. Browser-neutral session API; used unchanged in Pyodide and tests."""
import io
import json
from pathlib import Path
import tempfile
import zipfile

from .analysis import analyze_flight
from .export import label, ordered_codes, units, write_csv, write_xlsx
from .model import ConversionError
from .reader import read_jpi

MAX_INPUT_BYTES = 32 * 1024 * 1024


class BrowserSession:
    def __init__(self):
        self.downloads = {}
        self.next_id = 1
        self.analysis = {}

    def clear(self):
        self.downloads.clear()
        self.analysis.clear()

    def import_file(self, name, content):
        if len(content) > MAX_INPUT_BYTES:
            raise ConversionError('File exceeds the browser limit of 32 MiB')
        name = Path(name.replace('\\', '/')).name
        if not name.lower().endswith('.jpi'):
            raise ConversionError('Choose a .JPI file')
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / name
            path.write_bytes(content)
            download = read_jpi(path)
        # Never expose a temporary filesystem path as source identity.
        download.path = Path(name)
        for flight in download.flights:
            flight.source = download.path
        key = str(self.next_id)
        self.next_id += 1
        self.downloads[key] = download
        return {'key': key, 'name': name, 'metadata': download.metadata,
                'warnings': download.warnings, 'alarms': download.alarms,
                'flights': [self.flight_payload(key, i) for i in range(len(download.flights))]}

    def resolve(self, key):
        try:
            source, index = key.split(':')
            download = self.downloads[source]
            index = int(index)
            if index < 0:
                raise ValueError()
            return download.flights[index], download
        except (ValueError, KeyError, IndexError) as exc:
            raise ConversionError('Flight is no longer available; import its file again') from exc

    def flight_payload(self, source, index):
        flight, download = self.resolve(f'{source}:{index}')
        return {'key': f'{source}:{index}', 'id': flight.id, 'source': download.path.name,
                'start': flight.start.isoformat(sep=' '), 'duration': flight.duration,
                'samples': flight.samples, 'elapsed': flight.elapsed, 'series': flight.series,
                'channels': [{'code': code, 'label': label(code, flight), 'unit': units(code, download)}
                             for code in ordered_codes(flight) if code != 'Limits']}

    def analyze(self, key, minimum_seconds=10):
        flight, download = self.resolve(key)
        findings = analyze_flight(flight, download, minimum_seconds)
        result = {'minimumSeconds': minimum_seconds, 'findings': findings,
                  'scope': 'Recorded alarm limits and entirely missing channels only. '
                           'No flight-phase classification or mechanical diagnosis.'}
        self.analysis[key] = result
        return result

    def export(self, keys, mode='xlsx', graphs=True):
        if not keys or len(keys) != len(set(keys)):
            raise ConversionError('Select at least one distinct flight')
        if mode not in ('xlsx', 'csv', 'separate'):
            raise ConversionError('Unknown export format')
        selected = [self.resolve(key) for key in keys]
        downloads = list({id(d): d for _, d in selected}.values())
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            if mode == 'xlsx':
                path = root / 'JPI2Excel.xlsx'
                warnings = write_xlsx(path, selected, downloads, 0, graphs=graphs)
                self._append_findings(path, keys)
                return path.name, path.read_bytes(), warnings
            artifacts, warnings = [], []
            for ordinal, (key, (flight, download)) in enumerate(zip(keys, selected), 1):
                suffix = 'csv' if mode == 'csv' else 'xlsx'
                # Numeric prefix guarantees distinct names for duplicate IDs and sources.
                path = root / f'{ordinal}_Flight_{flight.id}.{suffix}'
                if mode == 'csv':
                    write_csv(path, flight, download)
                else:
                    warnings.extend(write_xlsx(path, [(flight, download)], [download], 0,
                                               summary=False, graphs=graphs))
                    self._append_findings(path, [key])
                artifacts.append(path)
            if len(artifacts) == 1:
                return artifacts[0].name, artifacts[0].read_bytes(), warnings
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
                for path in artifacts:
                    archive.writestr(path.name, path.read_bytes())
            return 'JPI2Excel.zip', buffer.getvalue(), warnings

    def _append_findings(self, path, keys):
        checked = [key for key in keys if key in self.analysis]
        if not checked:
            return
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        sheet = workbook.create_sheet('Findings')
        sheet.append(['Source', 'Flight', 'Sensor', 'Observation', 'Start (s)', 'End (s)',
                      'Span (s)', 'Extreme', 'Recorded limit', 'Unit', 'Rule', 'Minimum span (s)'])
        sheet.freeze_panes = 'A2'
        for key in checked:
            flight, download = self.resolve(key)
            result = self.analysis[key]
            sheet.append([download.path.name, flight.id, '', result['scope']])
            if not result['findings']:
                sheet.append([download.path.name, flight.id, '', 'No findings from these checks.'])
            for item in result['findings']:
                sheet.append([download.path.name, flight.id, item['label'], item['message'],
                              item['start'], item['end'], item['duration'], item['peak'],
                              item['threshold'], item['unit'], item['rule'], result['minimumSeconds']])
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        workbook.save(path)
        workbook.close()


def payload_json(value):
    return json.dumps(value, default=str, allow_nan=False)
