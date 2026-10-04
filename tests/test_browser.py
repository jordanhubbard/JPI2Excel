# Copyright © 2025-2026 by Alan M. Marcum, Nescorna Professional.
# All rights reserved.
"""Revision: 1. Browser session, export parity, and observed-span regressions."""
import csv
from datetime import datetime
import io
import json
from pathlib import Path
import unittest
import zipfile

from openpyxl import load_workbook
from jpi2excel.analysis import analyze_flight
from jpi2excel.browser import BrowserSession, MAX_INPUT_BYTES, payload_json
from jpi2excel.export import rows
from jpi2excel.model import ConversionError, Download, Flight
from jpi2excel.reader import read_jpi
from test_parser import synthetic_file

ROOT = Path(__file__).resolve().parents[1]


class AnalysisTests(unittest.TestCase):
    def analyze(self, values, elapsed=None, minimum=2, alarms=None):
        flight = Flight(1, Path('test.JPI'), datetime(2026, 1, 1), 2, 200, 0,
                        {'C1': values}, elapsed or list(range(0, len(values) * 2, 2)), [])
        download = Download(Path('test.JPI'), {'Engine temperature units': 'F'}, [], [flight],
                            alarms or {'CHT': (None, 400)})
        return analyze_flight(flight, download, minimum)

    def test_equality_extreme_and_duration(self):
        finding, = self.analyze([390, 400, 420, 410, 390])
        self.assertEqual((finding['start'], finding['end'], finding['duration'], finding['peak']), (2, 6, 4, 420))

    def test_missing_breaks_run_and_last_sample_adds_no_duration(self):
        self.assertEqual(self.analyze([410, None, 410]), [])
        finding, = self.analyze([410, 420])
        self.assertEqual(finding['duration'], 2)

    def test_variable_intervals_use_timestamps(self):
        finding, = self.analyze([400, 410, 420], elapsed=[0, 1, 10], minimum=10)
        self.assertEqual(finding['duration'], 10)

    def test_low_limit_and_missing_channels(self):
        finding, = self.analyze([90, 95, 96], alarms={'CHT': (95, None)})
        self.assertEqual((finding['kind'], finding['peak']), ('low limit', 90))
        finding, = self.analyze([None, None])
        self.assertEqual(finding['rule'], 'missing-channel-v1')
        self.assertEqual(self.analyze([390, 399]), [])

    def test_invalid_duration(self):
        for value in [-1, float('nan'), float('inf')]:
            with self.assertRaises(ConversionError):
                self.analyze([400, 410], minimum=value)


class BrowserTests(unittest.TestCase):
    def setUp(self):
        self.session = BrowserSession()

    def test_all_original_fixtures_match_native_reader(self):
        for path in sorted((ROOT / 'testdata').glob('*.JPI')):
            result = json.loads(payload_json(self.session.import_file(path.name, path.read_bytes())))
            native = read_jpi(path)
            self.assertEqual(len(result['flights']), len(native.flights))
            for actual, expected in zip(result['flights'], native.flights):
                self.assertEqual(actual['series'], expected.series)
                self.assertEqual(actual['elapsed'], expected.elapsed)
                self.assertEqual(actual['samples'], expected.samples)

    def test_corrupt_input_is_atomic_and_duplicates_remain_distinct(self):
        raw = synthetic_file()
        first = self.session.import_file('../same.JPI', raw)
        with self.assertRaises(ConversionError):
            self.session.import_file('bad.JPI', raw[:-1])
        second = self.session.import_file('same.JPI', raw)
        self.assertEqual(len(self.session.downloads), 2)
        self.assertEqual(first['name'], 'same.JPI')
        self.assertNotEqual(first['flights'][0]['key'], second['flights'][0]['key'])

    def test_csv_and_excel_preserve_values_and_graphs(self):
        result = self.session.import_file('sample.JPI', synthetic_file())
        key = result['flights'][0]['key']
        flight, download = self.session.resolve(key)
        name, data, _ = self.session.export([key], 'csv')
        parsed = list(csv.reader(io.StringIO(data.decode())))
        self.assertEqual(len(parsed) - 1, flight.samples)
        self.assertEqual(parsed[1][3], '' if next(rows(flight, download))[3] is None else str(next(rows(flight, download))[3]))
        self.session.analyze(key)
        name, data, _ = self.session.export([key], 'xlsx', graphs=True)
        book = load_workbook(io.BytesIO(data))
        self.assertIn('Findings', book.sheetnames)
        sheet = next(sheet for sheet in book if sheet.title.startswith('Flight'))
        self.assertEqual(list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))), [None if v == '' else v for v in next(rows(flight, download))])
        self.assertTrue(any(sheet._charts for sheet in book))
        book.close()

    def test_zip_separate_summary_and_clear(self):
        first = self.session.import_file('a.JPI', synthetic_file())['flights'][0]['key']
        second = self.session.import_file('a.JPI', synthetic_file())['flights'][0]['key']
        name, data, _ = self.session.export([first, second], 'separate', False)
        self.assertEqual(name, 'JPI2Excel.zip')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(len(set(archive.namelist())), 2)
            book = load_workbook(io.BytesIO(archive.read(archive.namelist()[0])))
            self.assertEqual(book['Summary'].sheet_state, 'hidden')
            book.close()
        self.session.clear()
        with self.assertRaises(ConversionError):
            self.session.resolve(first)

    def test_validation_and_limits(self):
        for name, data in [('bad.txt', b'anything'), ('huge.JPI', bytes(MAX_INPUT_BYTES + 1))]:
            with self.assertRaises(ConversionError):
                self.session.import_file(name, data)
        for keys, mode in [([], 'xlsx'), (['1:0', '1:0'], 'xlsx'), (['1:0'], 'bad')]:
            with self.assertRaises(ConversionError):
                self.session.export(keys, mode)
        with self.assertRaises(ConversionError):
            self.session.resolve('1:-1')
